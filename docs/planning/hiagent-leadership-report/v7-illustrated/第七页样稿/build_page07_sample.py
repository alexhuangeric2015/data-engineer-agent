# -*- coding: utf-8 -*-
"""按领导汇报页构图技能，单独生成第 07 页功能全景样稿。

素材只取自架构视觉版 V7 第 07 页已有标题、条目、页脚和备注。
本脚本不读取、不覆盖总 PPT，也不改 pages 与 content_layouts。
"""

from __future__ import annotations

import subprocess
import sys
from pathlib import Path

from lxml import etree
from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

DIR = Path(__file__).resolve().parent
PPTX_PATH = DIR / "第07页-功能全景-技能样稿.pptx"
PNG_PATH = DIR / "第07页-功能全景-技能样稿.png"
FONT_NAME = "Noto Sans CJK SC"
FONT_REGULAR = "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"
FONT_BOLD = "/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc"

SLIDE_W = 13.333333
SLIDE_H = 7.5

# 参考页的蓝白层级：浅蓝画布、白卡片、深蓝字、橙色只作标题短线。
BG = "F4F8FD"
WHITE = "FFFFFF"
INK = "16324A"
TITLE = "0E2A47"
BLUE = "1A6FDB"
BLUE_SOFT = "E8F2FC"
HERO = "EAF3FF"
HERO_LINE = "C5DCF6"
BORDER = "D5E4F2"
MUTED = "5C738C"
ORANGE = "F08C28"
CHIP = "F5F9FD"
PLAN = "EEF2F6"


def rgb(value: str) -> RGBColor:
    """把十六进制颜色转为 PPT 颜色。"""
    return RGBColor.from_string(value)


def emu(inches: float) -> Emu:
    """英寸转 EMU，供文字宽度比较使用。"""
    return Inches(inches)


def text_width(text: str, size: float, bold: bool) -> float:
    """用同一套中文字体测量单行宽度，单位为英寸。"""
    font = ImageFont.truetype(FONT_BOLD if bold else FONT_REGULAR, size=int(size * 8), index=0)
    box = font.getbbox(text)
    return (box[2] - box[0]) / 8 / 72


def add_shadow(shape) -> None:
    """给白卡片一层短阴影，让分组从浅蓝底上分开。"""
    effect = etree.SubElement(shape._element.spPr, qn("a:effectLst"))
    shadow = etree.SubElement(effect, qn("a:outerShdw"))
    shadow.set("blurRad", "63500")
    shadow.set("dist", "25400")
    shadow.set("dir", "2700000")
    shadow.set("algn", "tl")
    shadow.set("rotWithShape", "0")
    color = etree.SubElement(shadow, qn("a:srgbClr"))
    color.set("val", "8FB0D0")
    alpha = etree.SubElement(color, qn("a:alpha"))
    alpha.set("val", "28000")


def add_shape(slide, x, y, w, h, fill, line=None, radius=0.12, shadow=False, oval=False):
    """画圆角矩形或圆。radius 用形状自身的圆角比例。"""
    kind = MSO_SHAPE.OVAL if oval else MSO_SHAPE.ROUNDED_RECTANGLE
    shape = slide.shapes.add_shape(kind, emu(x), emu(y), emu(w), emu(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(0.75)
    else:
        shape.line.fill.background()
    if not oval:
        shape.adjustments[0] = radius
    if shadow:
        add_shadow(shape)
    return shape


def add_line(slide, x1, y1, x2, y2, color=BORDER) -> None:
    """画分组里的细分隔线。"""
    shape = slide.shapes.add_connector(1, emu(x1), emu(y1), emu(x2), emu(y2))
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(0.75)


def add_text(slide, x, y, w, h, text, size, color, bold=False, align="left", anchor="middle") -> None:
    """写单行或多行文字，并在超出文本框时直接失败。"""
    for line in text.split("\n"):
        if line and text_width(line, size, bold) > w + 0.02:
            raise ValueError(f"文字超出宽度 {w:.2f} 英寸：{line}")
    box = slide.shapes.add_textbox(emu(x), emu(y), emu(w), emu(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Emu(0)
    frame.margin_top = frame.margin_bottom = Emu(0)
    frame.vertical_anchor = {"middle": MSO_ANCHOR.MIDDLE, "top": MSO_ANCHOR.TOP}[anchor]
    aligns = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
    for index, line in enumerate(text.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.alignment = aligns[align]
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(0)
        run = paragraph.runs[0]
        run.font.name = FONT_NAME
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)
        latin = run._r.get_or_add_rPr()
        for tag in ("a:latin", "a:ea", "a:cs"):
            node = latin.find(qn(tag))
            if node is None:
                node = etree.SubElement(latin, qn(tag))
            node.set("typeface", FONT_NAME)


def add_pill(slide, x, y, w, h, text, fill, color, size=12, bold=True) -> None:
    """画状态或入口标签。"""
    add_shape(slide, x, y, w, h, fill, radius=0.5)
    add_text(slide, x, y, w, h, text, size, color, bold, "center")


def add_icon(slide, x, y, d, kind) -> None:
    """用能在小尺寸辨认的原生形状区分对象。"""
    add_shape(slide, x, y, d, d, BLUE_SOFT, oval=True)
    inner = d * 0.56
    left = x + (d - inner) / 2
    top = y + (d - inner) / 2
    if kind == "can":
        # 单缸比多层椭圆更不容易糊成色块。
        glyph = slide.shapes.add_shape(MSO_SHAPE.CAN, emu(left), emu(top), emu(inner), emu(inner))
        glyph.fill.solid()
        glyph.fill.fore_color.rgb = rgb(BLUE)
        glyph.line.fill.background()
    elif kind == "flow":
        glyph = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, emu(left), emu(top + inner * 0.28), emu(inner), emu(inner * 0.44))
        glyph.fill.solid()
        glyph.fill.fore_color.rgb = rgb(BLUE)
        glyph.line.fill.background()
    elif kind == "table":
        add_shape(slide, left, top, inner, inner, WHITE, BLUE, radius=0.08)
        add_line(slide, left, top + inner * 0.38, left + inner, top + inner * 0.38, BLUE)
        add_line(slide, left, top + inner * 0.68, left + inner, top + inner * 0.68, BLUE)
    elif kind == "book":
        add_shape(slide, left, top, inner, inner, WHITE, BLUE, radius=0.08)
        add_line(slide, left + inner * 0.5, top, left + inner * 0.5, top + inner, BLUE)
    elif kind == "link":
        radius = inner * 0.30
        mid = top + (inner - radius) / 2
        add_shape(slide, left, mid, radius, radius, WHITE, BLUE, oval=True)
        add_shape(slide, left + inner - radius, mid, radius, radius, WHITE, BLUE, oval=True)
        add_shape(slide, left + radius * 0.7, top + inner * 0.46, inner - radius * 1.4, inner * 0.08, BLUE, radius=0.5)
    elif kind == "check":
        add_shape(slide, left, top, inner, inner, WHITE, BLUE, radius=0.16)
        add_line(slide, left + inner * 0.24, top + inner * 0.52, left + inner * 0.44, top + inner * 0.72, BLUE)
        add_line(slide, left + inner * 0.44, top + inner * 0.72, left + inner * 0.78, top + inner * 0.28, BLUE)
    else:
        add_shape(slide, left + inner * 0.12, top, inner * 0.76, inner, WHITE, BLUE, radius=0.08)
        add_line(slide, left + inner * 0.28, top + inner * 0.32, left + inner * 0.72, top + inner * 0.32, BLUE)
        add_line(slide, left + inner * 0.28, top + inner * 0.55, left + inner * 0.72, top + inner * 0.55, BLUE)


def add_panel(slide, x, y, w, h):
    """画一列白卡片。"""
    add_shape(slide, x, y, w, h, WHITE, BORDER, radius=0.08, shadow=True)


def add_compact_title(slide, x, y, w, title, judgment) -> float:
    """短列头，给下方卡片留出高度。"""
    add_text(slide, x + 0.16, y + 0.08, w - 0.32, 0.30, title, 16, TITLE, True)
    add_text(slide, x + 0.16, y + 0.38, w - 0.32, 0.24, judgment, 12, BLUE)
    add_line(slide, x + 0.16, y + 0.66, x + w - 0.16, y + 0.66)
    return y + 0.76


def add_service_panel(slide, x, y, w, h) -> None:
    """归属是主结构：主体卡片最大，同级服务横向并排。"""
    add_panel(slide, x, y, w, h)
    cursor = add_compact_title(slide, x, y, w, "专业服务", "资产Agent为具体服务主体")
    hero_h = 1.58
    add_shape(slide, x + 0.16, cursor, w - 0.32, hero_h, HERO, HERO_LINE, radius=0.1)
    add_icon(slide, x + 0.32, cursor + 0.16, 0.42, "can")
    add_text(slide, x + 0.88, cursor + 0.14, 3.40, 0.36, "资产Agent", 20, TITLE, True)
    add_pill(slide, x + w - 1.08, cursor + 0.18, 0.74, 0.30, "当前", BLUE, WHITE, 13)
    add_text(slide, x + 0.88, cursor + 0.52, 5.40, 0.28, "表字段、字典码值为内部能力", 13, MUTED)
    chip_y = cursor + 0.96
    chip_w = (w - 0.64 - 0.14) / 2
    add_shape(slide, x + 0.32, chip_y, chip_w, 0.46, WHITE, HERO_LINE, radius=0.12)
    add_shape(slide, x + 0.32 + chip_w + 0.14, chip_y, chip_w, 0.46, WHITE, HERO_LINE, radius=0.12)
    add_text(slide, x + 0.32, chip_y, chip_w, 0.46, "表字段", 16, TITLE, True, "center")
    add_text(slide, x + 0.32 + chip_w + 0.14, chip_y, chip_w, 0.46, "字典码值", 16, TITLE, True, "center")

    label_y = cursor + hero_h + 0.10
    add_text(slide, x + 0.16, label_y, w - 0.32, 0.26, "与资产Agent同级", 13, MUTED, True)
    peers = [
        ("生产", "需求 / 设计 / 开发"),
        ("交付", "测试 / 评审"),
        ("运营分析", "治理 / 运维 / 分析"),
    ]
    card_y = label_y + 0.30
    card_h = y + h - 0.16 - card_y
    card_w = (w - 0.32 - 0.24) / 3
    for index, (name, detail) in enumerate(peers):
        card_x = x + 0.16 + index * (card_w + 0.12)
        add_shape(slide, card_x, card_y, card_w, card_h, CHIP, radius=0.1)
        add_text(slide, card_x + 0.12, card_y + 0.10, card_w - 0.86, 0.32, name, 15, TITLE, True)
        add_pill(slide, card_x + card_w - 0.74, card_y + 0.12, 0.60, 0.28, "规划", PLAN, MUTED, 12)
        add_text(slide, card_x + 0.12, card_y + 0.48, card_w - 0.24, 0.40, detail, 13, INK)


def add_content_grid(slide, x, y, w, h) -> None:
    """内容用四宫格，避免再排成和服务列一样的长条。"""
    add_panel(slide, x, y, w, h)
    cursor = add_compact_title(slide, x, y, w, "内容与方法", "SKILL提供方法，由Agent实际调用资源")
    items = ["术语 / 指标 / 模型语义", "规则 / 模板 / SKILL", "知识文件 / 检索索引", "事实查询数据 / 来源版本"]
    gap = 0.10
    tile_w = (w - 0.32 - gap) / 2
    tile_h = (y + h - 0.14 - cursor - gap) / 2
    for index, name in enumerate(items):
        tile_x = x + 0.16 + (index % 2) * (tile_w + gap)
        tile_y = cursor + (index // 2) * (tile_h + gap)
        add_shape(slide, tile_x, tile_y, tile_w, tile_h, CHIP, radius=0.1)
        add_text(slide, tile_x + 0.08, tile_y, tile_w - 0.16, tile_h, name, 13, TITLE, True, "center")


def add_platform_panel(slide, x, y, w, h) -> None:
    """连接、运营左右分列，每条独占一行，避免两行叠字。"""
    add_panel(slide, x, y, w, h)
    cursor = add_compact_title(slide, x, y, w, "平台连接与运营", "模块可复用现有系统承载")
    groups = [
        ("连接", ["目录 / 指标 / 血缘接口", "开发 / 调度 / 质量接入"]),
        ("运营", ["身份 / 权限 / 操作确认", "评测 / 反馈 / 内容更新"]),
    ]
    col_gap = 0.12
    col_w = (w - 0.32 - col_gap) / 2
    body_h = y + h - 0.14 - cursor
    line_h = (body_h - 0.34) / 2
    for index, (label, lines) in enumerate(groups):
        col_x = x + 0.16 + index * (col_w + col_gap)
        add_text(slide, col_x, cursor, col_w, 0.26, label, 13, BLUE, True)
        for line_index, line in enumerate(lines):
            line_y = cursor + 0.32 + line_index * (line_h + 0.06)
            add_shape(slide, col_x, line_y, col_w, line_h - 0.06, CHIP, radius=0.1)
            add_text(slide, col_x + 0.08, line_y, col_w - 0.16, line_h - 0.06, line, 13, INK, True, "center")


def add_entry(slide) -> None:
    """业务入口保持原文，右侧写明这一页不回答部署。"""
    add_shape(slide, 0.38, 1.00, 12.57, 0.48, WHITE, BORDER, radius=0.1, shadow=True)
    add_pill(slide, 0.52, 1.09, 1.18, 0.30, "业务入口", BLUE, WHITE, 12)
    entries = ["对话工作台", "数据平台入口", "按场景嵌入"]
    cursor = 1.86
    for name in entries:
        width = text_width(name, 13, True) + 0.36
        add_pill(slide, cursor, 1.09, width, 0.30, name, BLUE_SOFT, BLUE, 13)
        cursor += width + 0.12
    add_text(slide, 8.55, 1.00, 4.22, 0.48, "回答建什么，不表达服务器部署", 13, MUTED, align="right")


def add_landing(slide) -> None:
    """页脚是并列落点，不用箭头画成建设顺序。复用约束只留在平台列。"""
    add_shape(slide, 0.38, 5.96, 12.57, 1.08, WHITE, BORDER, radius=0.08, shadow=True)
    add_text(slide, 0.54, 6.08, 2.30, 0.40, "功能建设分别落实到", 13, TITLE, True)
    labels = ["Agent", "内容维护", "平台适配", "持续运营"]
    cursor = 3.00
    for label in labels:
        add_shape(slide, cursor, 6.12, 2.20, 0.36, BLUE_SOFT, radius=0.14)
        add_text(slide, cursor, 6.12, 2.20, 0.36, label, 13, BLUE, True, "center")
        cursor += 2.36
    add_text(
        slide,
        0.54,
        6.56,
        12.20,
        0.36,
        "逻辑模块不等于新增独立系统或已部署组件。",
        13,
        MUTED,
    )


def add_header(slide) -> None:
    """标题、短下划线和判断句分三层，避免互相压住。"""
    add_text(slide, 0.42, 0.06, 8.40, 0.42, "功能全景", 30, TITLE, True)
    add_shape(slide, 0.44, 0.50, 0.72, 0.05, ORANGE, radius=0.5)
    add_text(slide, 0.42, 0.60, 10.40, 0.32, "服务、内容、连接与运营形成完整建设范围", 15, INK)
    add_text(slide, 11.20, 0.12, 1.72, 0.30, "HiAgent", 14, BLUE, True, "right")


def add_footer(slide) -> None:
    """样稿仍标成第 07 页，避免被看成一份新的封面稿。"""
    add_text(slide, 0.42, 7.14, 8.60, 0.24, "企业数据智能体建设规划  /  规划视图  /  2026.10", 11, MUTED)
    add_text(slide, 11.40, 7.14, 1.52, 0.24, "07", 12, TITLE, True, "right")


def build() -> Presentation:
    """组装单页并写入备注中的原口径。"""
    presentation = Presentation()
    presentation.slide_width = emu(SLIDE_W)
    presentation.slide_height = emu(SLIDE_H)
    presentation.core_properties.title = "功能全景｜服务、内容、连接与运营形成完整建设范围"
    presentation.core_properties.subject = "第07页技能样稿，不替换架构视觉版V7正稿"
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(BG)

    add_header(slide)
    add_entry(slide)
    # 左侧归属区明显大于右侧两组，右侧内容与连接的内部结构也不相同。
    add_service_panel(slide, 0.34, 1.60, 7.56, 4.24)
    add_content_grid(slide, 8.04, 1.60, 4.94, 1.96)
    add_platform_panel(slide, 8.04, 3.68, 4.94, 2.16)
    add_landing(slide)
    add_footer(slide)
    notes = (
        "功能全景回答建什么，不表达服务器部署。专业服务层中资产Agent为具体服务主体，"
        "表字段、字典码值为其内部能力；其余Agent与资产Agent同级，避免把当前单Agent功能等同所有平台功能。\n\n"
        "架构口径：当前仅建设数据资产智能体；其他为目标规划。专业Agent内部规划并按需引用其他Agent。"
        "SKILL提供方法，由Agent实际调用资源。实际血缘经数据资产平台获取。"
        "图中逻辑模块不等于新增独立系统或已部署组件。"
    )
    slide.notes_slide.notes_text_frame.text = notes
    return presentation


def assert_source_text(presentation: Presentation) -> None:
    """样稿必须保留第 07 页原文中的功能条目。"""
    texts = []
    for shape in presentation.slides[0].shapes:
        if shape.has_text_frame:
            texts.append(shape.text_frame.text)
    joined = "\n".join(texts)
    required = [
        "功能全景",
        "服务、内容、连接与运营形成完整建设范围",
        "对话工作台",
        "数据平台入口",
        "按场景嵌入",
        "专业服务",
        "资产Agent",
        "当前",
        "表字段",
        "字典码值",
        "内部能力",
        "生产",
        "需求 / 设计 / 开发",
        "交付",
        "测试 / 评审",
        "运营分析",
        "治理 / 运维 / 分析",
        "规划",
        "同级",
        "内容与方法",
        "术语 / 指标 / 模型语义",
        "规则 / 模板 / SKILL",
        "知识文件 / 检索索引",
        "事实查询数据 / 来源版本",
        "平台连接与运营",
        "目录 / 指标 / 血缘接口",
        "开发 / 调度 / 质量接入",
        "身份 / 权限 / 操作确认",
        "评测 / 反馈 / 内容更新",
        "内容维护",
        "平台适配",
        "持续运营",
        "模块可复用现有系统承载",
        "不表达服务器部署",
        "不等于新增独立系统或已部署组件",
    ]
    missing = [item for item in required if item not in joined]
    if missing:
        raise SystemExit("缺少原文：" + "、".join(missing))


def export_png() -> None:
    """用演示文稿导出预览，避免另画一套看图。"""
    work = Path("/tmp/page07-sample")
    work.mkdir(exist_ok=True)
    profile = work / "lo-profile"
    profile.mkdir(exist_ok=True)
    subprocess.run(
        [
            "soffice",
            "--headless",
            "--norestore",
            "--nolockcheck",
            f"-env:UserInstallation=file://{profile}",
            "--convert-to",
            "pdf",
            "--outdir",
            str(work),
            str(PPTX_PATH),
        ],
        check=True,
        capture_output=True,
        text=True,
    )
    pdf = work / (PPTX_PATH.stem + ".pdf")
    prefix = work / "preview"
    subprocess.run(["pdftoppm", "-png", "-r", "144", "-f", "1", "-l", "1", str(pdf), str(prefix)], check=True)
    produced = prefix.with_name(prefix.name + "-1.png")
    produced.replace(PNG_PATH)


def main() -> None:
    """生成 pptx 与 png。"""
    presentation = build()
    assert_source_text(presentation)
    presentation.save(PPTX_PATH)
    export_png()
    print(PPTX_PATH)
    print(PNG_PATH)


if __name__ == "__main__":
    try:
        main()
    except subprocess.CalledProcessError as exc:
        sys.stderr.write(exc.stderr or exc.stdout or str(exc))
        raise
