from pathlib import Path
import json, math
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement

ROOT=Path(__file__).parent
(ROOT/'qa').mkdir(exist_ok=True)
OUT=ROOT/'HiAgent企业数据智能体体系规划_蓝白咨询版_V3.pptx'
FONT='Noto Sans CJK SC'
BLUE='266DE7'; TITLE='1450B6'; DEEP='01379C'; LIGHT='A8C5F5'; PALE='EAF1FF'; SOFT='F5F8FD'
INK='202B3C'; GRAY='526176'; BORDER='D2DFF3'; RED='C00000'; WHITE='FFFFFF'
prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
prs.core_properties.title='HiAgent 企业数据智能体体系规划｜蓝白咨询版 V3'
prs.core_properties.subject='总体蓝图、专业能力、已有资产治理与协同演进'
prs.core_properties.author='项目规划组'
prs.core_properties.keywords='HiAgent,企业数据智能体,上下文工程,Agent Skills,MCP,AgentOps'
records=[];manifest=[];notes_out=[]
REFS={
 'R1':('LangGraph / Deep Agents','https://github.com/langchain-ai/deepagents'),
 'R2':('Agent Skills','https://github.com/anthropics/skills'),
 'R3':('上下文工程','https://www.anthropic.com/engineering/effective-context-engineering-for-ai-agents'),
 'R4':('MCP','https://github.com/modelcontextprotocol/servers'),
 'R5':('DataHub Agent Context Kit','https://github.com/datahub-project/datahub/blob/master/docs/dev-guides/agent-context/agent-context.md'),
 'R6':('GraphRAG','https://github.com/microsoft/graphrag'),
 'R7':('Langfuse','https://github.com/langfuse/langfuse')}

def rgb(v):return RGBColor.from_string(v)
def clean(a):
    a._element.spPr.append(OxmlElement('a:effectLst'))
    for e in a._element.xpath('./p:style/a:effectRef'):e.set('idx','0')
def box(s,x,y,w,h,fill=WHITE,stroke=None,r=0,dash=False):
    a=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    a.fill.solid();a.fill.fore_color.rgb=rgb(fill)
    if stroke:
        a.line.color.rgb=rgb(stroke);a.line.width=Pt(.65)
        if dash:
            v=OxmlElement('a:prstDash');v.set('val','dash');a._element.spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ln').append(v)
    else:a.line.fill.background()
    if r:a.adjustments[0]=r
    clean(a);return a
def gradient(s,x,y,w,h,left=PALE,right=LIGHT):
    a=box(s,x,y,w,h)
    sp=a._element.spPr
    for f in list(sp):
        if f.tag.endswith('solidFill'):sp.remove(f)
    g=OxmlElement('a:gradFill');ls=OxmlElement('a:gsLst')
    for pos,col in [('0',left),('100000',right)]:
        st=OxmlElement('a:gs');st.set('pos',pos);co=OxmlElement('a:srgbClr');co.set('val',col);st.append(co);ls.append(st)
    g.append(ls);ln=OxmlElement('a:lin');ln.set('ang','0');ln.set('scaled','1');g.append(ln);sp.append(g);return a
def tx(s,x,y,w,h,t,size=15,c=INK,bold=False,align='left',top=False,link=None):
    a=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=a.text_frame
    tf.clear();tf.word_wrap=True;tf.margin_top=tf.margin_bottom=tf.margin_left=tf.margin_right=0
    tf.vertical_anchor=MSO_ANCHOR.TOP if top else MSO_ANCHOR.MIDDLE
    for i,line in enumerate(t.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line
        p.font.name=FONT;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(c)
        p.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[align]
        p.line_spacing=1.12;p.space_after=Pt(2);p.space_before=Pt(0)
        rpr=p._p.get_or_add_pPr().get_or_add_defRPr();e=OxmlElement('a:ea');e.set('typeface',FONT);rpr.append(e)
        if link and p.runs:p.runs[0].hyperlink.address=link
    records.append(dict(slide=len(prs.slides),x=x,y=y,w=w,h=h,text=t,size=size))
    return a
def line(s,x1,y1,x2,y2,c=BLUE,width=.9,dash=False):
    a=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));a.line.color.rgb=rgb(c);a.line.width=Pt(width);clean(a)
    if dash:
        e=OxmlElement('a:prstDash');e.set('val','dash');a._element.spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ln').append(e)
    return a
def arrow(s,x1,y1,x2,y2,c=BLUE,width=1.2):
    a=line(s,x1,y1,x2,y2,c,width);angle=math.atan2(y2-y1,x2-x1)
    for delta in [-.55,.55]:line(s,x2,y2,x2-.075*math.cos(angle+delta),y2-.075*math.sin(angle+delta),c,width)
    return a
def circle(s,x,y,d,fill=BLUE,stroke=None):
    a=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(d),Inches(d));a.fill.solid();a.fill.fore_color.rgb=rgb(fill)
    if stroke:a.line.color.rgb=rgb(stroke)
    else:a.line.fill.background()
    clean(a);return a
def header(s,title,section):
    tx(s,.34,.17,10.9,.49,title,23,TITLE,True)
    tx(s,11.36,.18,1.62,.24,'HiAgent 数据智能体',9,DEEP,True,'right')
    tx(s,11.36,.46,1.62,.18,'体系规划 · V3',8,GRAY,align='right')
    line(s,.34,.78,12.98,.78,BORDER,.7)
    line(s,.34,7.22,12.98,7.22,BORDER,.65)
    tx(s,.35,7.29,11.1,.13,section+'  ｜  企业数据智能体体系规划  ｜  2026年10月',7.5,GRAY)
    tx(s,12.37,7.27,.61,.18,f'{len(prs.slides):02d}',9,GRAY,align='right')
def new(title,section,summary=None):
    s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb(WHITE)
    header(s,title,section)
    if summary:
        box(s,.34,.95,12.64,.59,PALE)
        tx(s,.48,1.01,12.36,.46,summary,14.5,TITLE,True)
    manifest.append(dict(page=len(prs.slides),title=title,section=section))
    return s
def note(s,body,refs=(),q=None):
    t=body+('\n\n领导追问：\n'+q if q else '')
    if refs:t+='\n\n一手参考：\n'+'\n'.join(k+' '+REFS[k][0]+' '+REFS[k][1] for k in refs)
    t+='\n\n统一口径：目标规划与当前建设分别表述；专业Agent内部规划并按需引用其他Agent；血缘事实来自数据资产平台。'
    s.notes_slide.notes_text_frame.text=t;notes_out.append(dict(page=len(prs.slides),title=manifest[-1]['title'],notes=t))
def bar(s,x,y,w,h,t,fill=BLUE,size=16):
    box(s,x,y,w,h,fill,r=.025);tx(s,x+.08,y+.03,w-.16,h-.06,t,size,WHITE,True,'center')
def panel(s,x,y,w,h,title,fill=WHITE,titlefill=BLUE):
    a=box(s,x,y,w,h,fill,BORDER,r=.025);bar(s,x+.01,y+.01,w-.02,.39,title,titlefill);return a
def chip(s,x,y,w,h,t,fill=PALE,c=DEEP,size=14,bold=False):
    box(s,x,y,w,h,fill,BORDER,r=.025);tx(s,x+.09,y+.04,w-.18,h-.08,t,size,c,bold,'center')
def footer_band(s,t,y=6.65):
    gradient(s,.34,y,12.64,.39,'EDF3FF','D8E6FF');tx(s,.48,y+.03,12.36,.32,t,13.5,DEEP,True,'center')
def table(s,x,y,widths,rows,rh=.56,size=13):
    for ri,row in enumerate(rows):
        xx=x
        for w,t in zip(widths,row):
            box(s,xx,y+ri*rh,w,rh,BLUE if ri==0 else (WHITE if ri%2 else SOFT),BORDER)
            tx(s,xx+.11,y+ri*rh+.05,w-.22,rh-.10,t,size,WHITE if ri==0 else INK,ri==0)
            xx+=w
def bullets(s,x,y,w,items,size=14,step=.45,c=INK):
    for i,t in enumerate(items):
        circle(s,x,y+i*step+.12,.045,BLUE)
        tx(s,x+.13,y+i*step,w-.13,step-.02,t,size,c,top=True)
