from pathlib import Path
import json,re
import fitz
from pptx import Presentation
from PIL import Image,ImageOps,ImageDraw,ImageFont
R=Path(__file__).parent;Q=R/'qa'
ppt=next(R.glob('*.pptx'));pdf=next(Q.glob('*.pdf'))
p=Presentation(ppt);d=fitz.open(pdf)
missing=[];bounds=[]
norm=lambda t:re.sub(r'\s+','',t).replace('→','').replace('·','')
for r in json.loads((Q/'textboxes.json').read_text()):
    area=fitz.Rect(r['x']*72-2,r['y']*72-2,(r['x']+r['w'])*72+2,(r['y']+r['h'])*72+2)
    found=d[r['slide']-1].get_text(clip=area)
    if norm(r['text']) not in norm(found):missing.append(dict(page=r['slide'],text=r['text'],found=found,rect=list(area)))
for i,s in enumerate(p.slides,1):
    for a in s.shapes:
        if min(a.left,a.top)<-10000 or a.left+a.width>p.slide_width+10000 or a.top+a.height>p.slide_height+10000:bounds.append([i,a.shape_id])
    pix=d[i-1].get_pixmap(matrix=fitz.Matrix(1.5,1.5),alpha=False);pix.save(Q/f'page-{i:02}.png')
full='\n'.join(a.text for s in p.slides for a in s.shapes if a.has_text_frame)
forbidden=[x for x in ['预算','测算','报价','技术备答','85%','95%','长沙银行','中电金信'] if x in full]
report=dict(slides=len(p.slides),pdf_pages=len(d),missing_text=missing,bounds=bounds,forbidden_visible=forbidden,notes=sum(bool(s.notes_slide.notes_text_frame.text) for s in p.slides))
(Q/'validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
font=ImageFont.truetype('/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc',19)
for start in range(1,len(d)+1,6):
    out=Image.new('RGB',(1600,3*477),'#dce6f3');dr=ImageDraw.Draw(out)
    for k,i in enumerate(range(start,min(start+6,len(d)+1))):
        im=Image.open(Q/f'page-{i:02}.png');im.thumbnail((780,439));x=10+(k%2)*800;y=10+(k//2)*477;out.paste(im,(x,y));dr.text((x,y+441),f'{i:02d}',font=font,fill='#1450b6')
    out.save(Q/f'contact-{start:02d}.png')
print(json.dumps(report,ensure_ascii=False)[:18000])
